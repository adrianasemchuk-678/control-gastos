"""
Módulo de generación de reporte Excel interactivo para Control de Gastos.
Produce un archivo .xlsx completo, funcional de control financiero, con fórmulas dinámicas,
estética pastel clara, tarjetas KPI, distribución por categoría, gráficos nativos de Excel y formato profesional.
"""

import io
from datetime import date
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, Reference
from openpyxl.utils import get_column_letter

import database as db
import categories as cat

FONT_FAMILY = "Segoe UI"

# Estilos de Fuentes
FONT_TITLE = Font(name=FONT_FAMILY, size=15, bold=True, color="BE123C")
FONT_SUBTITLE = Font(name=FONT_FAMILY, size=9.5, italic=True, color="64748B")
FONT_SECTION = Font(name=FONT_FAMILY, size=11, bold=True, color="4A2D3B")
FONT_HEADER = Font(name=FONT_FAMILY, size=10.5, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name=FONT_FAMILY, size=10, color="1E293B")
FONT_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="1E293B")
FONT_TOTAL = Font(name=FONT_FAMILY, size=11, bold=True, color="BE123C")

# Estilos de Relleno (Colores Pastel)
FILL_BANNER = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
FILL_HEADER_PRIMARY = PatternFill(start_color="FF6584", end_color="FF6584", fill_type="solid")
FILL_HEADER_ALT = PatternFill(start_color="FDA4AF", end_color="FDA4AF", fill_type="solid")
FILL_HEADER_DARK = PatternFill(start_color="BE123C", end_color="BE123C", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="FFF7F9", end_color="FFF7F9", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")

# Tarjetas KPI
FILL_KPI_INGRESO = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
FONT_KPI_INGRESO_TITLE = Font(name=FONT_FAMILY, size=8.5, bold=True, color="0369A1")
FONT_KPI_INGRESO_VAL = Font(name=FONT_FAMILY, size=15, bold=True, color="0369A1")

FILL_KPI_GASTO = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
FONT_KPI_GASTO_TITLE = Font(name=FONT_FAMILY, size=8.5, bold=True, color="BE123C")
FONT_KPI_GASTO_VAL = Font(name=FONT_FAMILY, size=15, bold=True, color="BE123C")

FILL_KPI_SALDO = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
FONT_KPI_SALDO_TITLE = Font(name=FONT_FAMILY, size=8.5, bold=True, color="15803D")
FONT_KPI_SALDO_VAL = Font(name=FONT_FAMILY, size=15, bold=True, color="15803D")

FILL_KPI_PCT = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
FONT_KPI_PCT_TITLE = Font(name=FONT_FAMILY, size=8.5, bold=True, color="B45309")
FONT_KPI_PCT_VAL = Font(name=FONT_FAMILY, size=15, bold=True, color="B45309")

# Estados de Pago
STATUS_STYLES = {
    "PAGADO": {
        "fill": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "font": Font(name=FONT_FAMILY, size=9.5, bold=True, color="166534"),
    },
    "VENCIDO": {
        "fill": PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid"),
        "font": Font(name=FONT_FAMILY, size=9.5, bold=True, color="9F1239"),
    },
    "POR VENCER": {
        "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "font": Font(name=FONT_FAMILY, size=9.5, bold=True, color="92400E"),
    },
    "PRÓXIMO": {
        "fill": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
        "font": Font(name=FONT_FAMILY, size=9.5, bold=True, color="475569"),
    }
}

# Bordes
THIN_SIDE = Side(style="thin", color="CBD5E1")
BORDER_CELL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_HEADER = Border(
    left=Side(style="thin", color="FDA4AF"),
    right=Side(style="thin", color="FDA4AF"),
    top=Side(style="thin", color="FDA4AF"),
    bottom=Side(style="medium", color="BE123C")
)
BORDER_TOTAL = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=Side(style="thin", color="BE123C"),
    bottom=Side(style="double", color="BE123C")
)

# Alineaciones
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# Formatos numéricos de Excel
FORMAT_CURRENCY = "$ #,##0.00"
FORMAT_PERCENT = "0.0%"
FORMAT_INTEGER = "#,##0"


def generar_excel_control(mes: int, anio: int, sueldo: float = 0.0) -> bytes:
    """
    Genera en memoria un archivo Excel (.xlsx) interactivo y completamente formateado.
    Contiene 3 pestañas:
    1. 📊 Panel de Control (Dashboard con tarjetas KPI, fórmulas dinámicas y gráficos).
    2. 📝 Gastos del Mes (Detalle pormenorizado con autofiltros y totales automáticos).
    3. 🔔 Pagos Fijos (Control de fechas de vencimiento, alertas y estados).
    """
    wb = openpyxl.Workbook()
    
    # 1. Crear Hojas
    ws_dash = wb.active
    ws_dash.title = "📊 Panel de Control"
    ws_gastos = wb.create_sheet("📝 Gastos del Mes")
    ws_pagos = wb.create_sheet("🔔 Pagos Fijos")
    
    for ws in (ws_dash, ws_gastos, ws_pagos):
        ws.views.sheetView[0].showGridLines = True
        
    nombres_meses = [
        "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
        "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
    ]
    mes_nombre = nombres_meses[mes - 1]
    
    # Obtener datos de la base de datos
    gastos = db.get_expenses(mes=mes, anio=anio)
    pagos_fijos = db.get_estado_pagos_fijos(mes=mes, anio=anio)
    
    # =========================================================================
    # HOJA 2: DETALLE DE GASTOS
    # =========================================================================
    _construir_hoja_gastos(ws_gastos, gastos, mes_nombre, anio)
    
    # =========================================================================
    # HOJA 3: PAGOS FIJOS
    # =========================================================================
    _construir_hoja_pagos_fijos(ws_pagos, pagos_fijos, mes_nombre, anio)
    
    # =========================================================================
    # HOJA 1: PANEL DE CONTROL (DASHBOARD)
    # =========================================================================
    last_gasto_row = max(len(gastos) + 3, 4)
    total_gasto_row = last_gasto_row + 1 if len(gastos) > 0 else 5
    
    last_pf_row = max(len(pagos_fijos) + 3, 4)
    total_pf_row = last_pf_row + 1 if len(pagos_fijos) > 0 else 5
    
    _construir_hoja_dashboard(
        ws_dash=ws_dash,
        sueldo=sueldo,
        mes_nombre=mes_nombre,
        anio=anio,
        gastos=gastos,
        last_gasto_row=last_gasto_row,
        total_gasto_row=total_gasto_row,
        last_pf_row=last_pf_row,
        total_pf_row=total_pf_row
    )
    
    # Autoajuste de columnas en las 3 hojas
    for ws in (ws_dash, ws_gastos, ws_pagos):
        _autoajustar_columnas(ws)
        
    # Guardar en memoria y retornar bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _construir_hoja_gastos(ws, gastos: List[Dict[str, Any]], mes_nombre: str, anio: int):
    # Banner
    ws.merge_cells("A1:E1")
    c_banner = ws["A1"]
    c_banner.value = f"📝 DETALLE COMPLETO DE GASTOS - {mes_nombre.upper()} {anio}"
    c_banner.font = FONT_TITLE
    c_banner.fill = FILL_BANNER
    c_banner.alignment = ALIGN_CENTER
    ws.row_dimensions[1].height = 36
    
    ws.merge_cells("A2:E2")
    c_sub = ws["A2"]
    c_sub.value = f"Total de movimientos registrados en el período: {len(gastos)} | Sistema de Control de Gastos"
    c_sub.font = FONT_SUBTITLE
    c_sub.alignment = ALIGN_CENTER
    ws.row_dimensions[2].height = 18
    
    # Encabezados
    headers = ["N°", "Fecha", "Descripción / Concepto", "Categoría", "Monto ($)"]
    ws.row_dimensions[3].height = 26
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER_PRIMARY
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_HEADER

    # Filas de datos
    start_row = 4
    if gastos:
        for idx, g in enumerate(gastos, start=1):
            curr_row = start_row + idx - 1
            ws.row_dimensions[curr_row].height = 20
            fill = FILL_ZEBRA if idx % 2 == 0 else FILL_WHITE
            
            c_num = ws.cell(row=curr_row, column=1, value=idx)
            c_num.alignment = ALIGN_CENTER
            
            c_fec = ws.cell(row=curr_row, column=2, value=str(g.get("fecha", "")))
            c_fec.alignment = ALIGN_CENTER
            
            c_desc = ws.cell(row=curr_row, column=3, value=str(g.get("descripcion", "")))
            c_desc.alignment = ALIGN_LEFT
            
            cat_nombre = str(g.get("categoria", ""))
            c_cat = ws.cell(row=curr_row, column=4, value=cat_nombre)
            c_cat.alignment = ALIGN_LEFT
            
            c_monto = ws.cell(row=curr_row, column=5, value=float(g.get("monto", 0.0)))
            c_monto.alignment = ALIGN_RIGHT
            c_monto.number_format = FORMAT_CURRENCY
            
            for c in (c_num, c_fec, c_desc, c_cat, c_monto):
                c.font = FONT_REGULAR
                c.fill = fill
                c.border = BORDER_CELL

        last_data_row = start_row + len(gastos) - 1
        total_row = last_data_row + 1
        
        # Fila de Total
        ws.row_dimensions[total_row].height = 24
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
        c_tot_lbl = ws.cell(row=total_row, column=1, value="TOTAL DE GASTOS ANOTADOS:")
        c_tot_lbl.font = FONT_TOTAL
        c_tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_lbl.fill = FILL_TOTAL
        
        c_tot_val = ws.cell(row=total_row, column=5, value=f"=SUM(E{start_row}:E{last_data_row})")
        c_tot_val.font = FONT_TOTAL
        c_tot_val.alignment = ALIGN_RIGHT
        c_tot_val.number_format = FORMAT_CURRENCY
        c_tot_val.fill = FILL_TOTAL
        
        for col_idx in range(1, 6):
            ws.cell(row=total_row, column=col_idx).border = BORDER_TOTAL
            
        # Autofiltro
        ws.auto_filter.ref = f"A3:E{last_data_row}"
    else:
        # Fila vacía informativa
        ws.merge_cells("A4:E4")
        c_empty = ws["A4"]
        c_empty.value = "No se registraron gastos en este mes."
        c_empty.font = FONT_SUBTITLE
        c_empty.alignment = ALIGN_CENTER
        c_empty.border = BORDER_CELL
        
        total_row = 5
        ws.merge_cells("A5:D5")
        ws["A5"] = "TOTAL DE GASTOS ANOTADOS:"
        ws["A5"].font = FONT_TOTAL
        ws["A5"].alignment = Alignment(horizontal="right", vertical="center")
        ws["A5"].fill = FILL_TOTAL
        ws["E5"] = 0.0
        ws["E5"].font = FONT_TOTAL
        ws["E5"].alignment = ALIGN_RIGHT
        ws["E5"].number_format = FORMAT_CURRENCY
        ws["E5"].fill = FILL_TOTAL
        for col_idx in range(1, 6):
            ws.cell(row=5, column=col_idx).border = BORDER_TOTAL


def _construir_hoja_pagos_fijos(ws, pagos_fijos: List[Dict[str, Any]], mes_nombre: str, anio: int):
    # Banner
    ws.merge_cells("A1:G1")
    c_banner = ws["A1"]
    c_banner.value = f"🔔 CONTROL DE PAGOS FIJOS Y VENCIMIENTOS - {mes_nombre.upper()} {anio}"
    c_banner.font = FONT_TITLE
    c_banner.fill = FILL_BANNER
    c_banner.alignment = ALIGN_CENTER
    ws.row_dimensions[1].height = 36
    
    ws.merge_cells("A2:G2")
    c_sub = ws["A2"]
    c_sub.value = "Servicios mensuales recurrentes, recordatorios y estado de cumplimiento"
    c_sub.font = FONT_SUBTITLE
    c_sub.alignment = ALIGN_CENTER
    ws.row_dimensions[2].height = 18
    
    headers = ["N°", "Servicio / Concepto", "Monto Estimado ($)", "Categoría", "Ventana de Pago", "Estado Actual", "Fecha de Pago"]
    ws.row_dimensions[3].height = 26
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER_DARK
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_HEADER

    start_row = 4
    if pagos_fijos:
        for idx, p in enumerate(pagos_fijos, start=1):
            curr_row = start_row + idx - 1
            ws.row_dimensions[curr_row].height = 20
            fill_base = FILL_ZEBRA if idx % 2 == 0 else FILL_WHITE
            
            c_num = ws.cell(row=curr_row, column=1, value=idx)
            c_num.alignment = ALIGN_CENTER
            
            c_nom = ws.cell(row=curr_row, column=2, value=str(p.get("nombre", "")))
            c_nom.alignment = ALIGN_LEFT
            
            c_monto = ws.cell(row=curr_row, column=3, value=float(p.get("monto", 0.0)))
            c_monto.alignment = ALIGN_RIGHT
            c_monto.number_format = FORMAT_CURRENCY
            
            c_cat = ws.cell(row=curr_row, column=4, value=str(p.get("categoria", "")))
            c_cat.alignment = ALIGN_LEFT
            
            d_desde = p.get("dia_desde", 1)
            d_hasta = p.get("dia_hasta", 10)
            c_vent = ws.cell(row=curr_row, column=5, value=f"Día {d_desde} al {d_hasta}")
            c_vent.alignment = ALIGN_CENTER
            
            # Estado formateado y coloreado
            estado_raw = str(p.get("estado", "proximo")).lower()
            if p.get("pagado"):
                estado_txt = "PAGADO"
            elif estado_raw == "vencido":
                estado_txt = "VENCIDO"
            elif estado_raw == "por_vencer":
                estado_txt = "POR VENCER"
            else:
                estado_txt = "PRÓXIMO"
                
            c_est = ws.cell(row=curr_row, column=6, value=estado_txt)
            c_est.alignment = ALIGN_CENTER
            st_style = STATUS_STYLES.get(estado_txt, STATUS_STYLES["PRÓXIMO"])
            c_est.font = st_style["font"]
            c_est.fill = st_style["fill"]
            c_est.border = BORDER_CELL
            
            fecha_pago = p.get("fecha_pago") or "-"
            c_fp = ws.cell(row=curr_row, column=7, value=str(fecha_pago))
            c_fp.alignment = ALIGN_CENTER
            
            for c in (c_num, c_nom, c_monto, c_cat, c_vent, c_fp):
                c.font = FONT_REGULAR
                c.fill = fill_base
                c.border = BORDER_CELL

        last_data_row = start_row + len(pagos_fijos) - 1
        total_row = last_data_row + 1
        
        # Fila de Total
        ws.row_dimensions[total_row].height = 24
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
        c_tot_lbl = ws.cell(row=total_row, column=1, value="TOTAL COMPROMISO FIJO:")
        c_tot_lbl.font = FONT_TOTAL
        c_tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_lbl.fill = FILL_TOTAL
        
        c_tot_val = ws.cell(row=total_row, column=3, value=f"=SUM(C{start_row}:C{last_data_row})")
        c_tot_val.font = FONT_TOTAL
        c_tot_val.alignment = ALIGN_RIGHT
        c_tot_val.number_format = FORMAT_CURRENCY
        c_tot_val.fill = FILL_TOTAL
        
        for col_idx in range(1, 8):
            ws.cell(row=total_row, column=col_idx).border = BORDER_TOTAL
            
        ws.auto_filter.ref = f"A3:G{last_data_row}"
    else:
        ws.merge_cells("A4:G4")
        c_empty = ws["A4"]
        c_empty.value = "No hay pagos fijos registrados en este período."
        c_empty.font = FONT_SUBTITLE
        c_empty.alignment = ALIGN_CENTER
        c_empty.border = BORDER_CELL
        
        total_row = 5
        ws.merge_cells("A5:B5")
        ws["A5"] = "TOTAL COMPROMISO FIJO:"
        ws["A5"].font = FONT_TOTAL
        ws["A5"].alignment = Alignment(horizontal="right", vertical="center")
        ws["A5"].fill = FILL_TOTAL
        ws["C5"] = 0.0
        ws["C5"].font = FONT_TOTAL
        ws["C5"].alignment = ALIGN_RIGHT
        ws["C5"].number_format = FORMAT_CURRENCY
        ws["C5"].fill = FILL_TOTAL
        for col_idx in range(1, 8):
            ws.cell(row=5, column=col_idx).border = BORDER_TOTAL


def _construir_hoja_dashboard(ws_dash, sueldo: float, mes_nombre: str, anio: int, gastos: List[Dict[str, Any]], last_gasto_row: int, total_gasto_row: int, last_pf_row: int, total_pf_row: int):
    # Banner Principal
    ws_dash.merge_cells("A1:I1")
    c_banner = ws_dash["A1"]
    c_banner.value = f"🌸 CONTROL FINANCIERO MENSUAL - ADRIANA ({mes_nombre.upper()} {anio})"
    c_banner.font = FONT_TITLE
    c_banner.fill = FILL_BANNER
    c_banner.alignment = ALIGN_CENTER
    ws_dash.row_dimensions[1].height = 36
    
    ws_dash.merge_cells("A2:I2")
    c_sub = ws_dash["A2"]
    c_sub.value = f"Panel de Control Ejecutivo y Seguimiento del Presupuesto | Generado: {date.today().strftime('%d/%m/%Y')}"
    c_sub.font = FONT_SUBTITLE
    c_sub.alignment = ALIGN_CENTER
    ws_dash.row_dimensions[2].height = 18
    
    ws_dash.row_dimensions[3].height = 10
    
    # ==========================================
    # TARJETAS KPI (FILAS 4 Y 5)
    # ==========================================
    ws_dash.row_dimensions[4].height = 18
    ws_dash.row_dimensions[5].height = 30
    
    # KPI 1: Ingreso
    ws_dash.merge_cells("B4:C4")
    ws_dash["B4"] = "💵 DINERO INICIAL / SUELDO"
    ws_dash["B4"].font = FONT_KPI_INGRESO_TITLE
    ws_dash["B4"].alignment = ALIGN_CENTER
    ws_dash["B4"].fill = FILL_KPI_INGRESO
    
    ws_dash.merge_cells("B5:C5")
    ws_dash["B5"] = float(sueldo)
    ws_dash["B5"].font = FONT_KPI_INGRESO_VAL
    ws_dash["B5"].alignment = ALIGN_CENTER
    ws_dash["B5"].fill = FILL_KPI_INGRESO
    ws_dash["B5"].number_format = FORMAT_CURRENCY

    # KPI 2: Total Gastado (Fórmula referenciando Hoja de Gastos)
    ws_dash.merge_cells("D4:E4")
    ws_dash["D4"] = "💸 LO QUE YA GASTASTE"
    ws_dash["D4"].font = FONT_KPI_GASTO_TITLE
    ws_dash["D4"].alignment = ALIGN_CENTER
    ws_dash["D4"].fill = FILL_KPI_GASTO
    
    ws_dash.merge_cells("D5:E5")
    ws_dash["D5"] = f"='📝 Gastos del Mes'!E{total_gasto_row}"
    ws_dash["D5"].font = FONT_KPI_GASTO_VAL
    ws_dash["D5"].alignment = ALIGN_CENTER
    ws_dash["D5"].fill = FILL_KPI_GASTO
    ws_dash["D5"].number_format = FORMAT_CURRENCY

    # KPI 3: Saldo Disponible (Fórmula = Ingreso - Gastos)
    ws_dash.merge_cells("F4:G4")
    ws_dash["F4"] = "🐷 DINERO DISPONIBLE"
    ws_dash["F4"].font = FONT_KPI_SALDO_TITLE
    ws_dash["F4"].alignment = ALIGN_CENTER
    ws_dash["F4"].fill = FILL_KPI_SALDO
    
    ws_dash.merge_cells("F5:G5")
    ws_dash["F5"] = "=B5-D5"
    ws_dash["F5"].font = FONT_KPI_SALDO_VAL
    ws_dash["F5"].alignment = ALIGN_CENTER
    ws_dash["F5"].fill = FILL_KPI_SALDO
    ws_dash["F5"].number_format = FORMAT_CURRENCY

    # KPI 4: % Consumido
    ws_dash.merge_cells("H4:I4")
    ws_dash["H4"] = "📊 % DEL SUELDO GASTADO"
    ws_dash["H4"].font = FONT_KPI_PCT_TITLE
    ws_dash["H4"].alignment = ALIGN_CENTER
    ws_dash["H4"].fill = FILL_KPI_PCT
    
    ws_dash.merge_cells("H5:I5")
    ws_dash["H5"] = "=IF(B5>0, D5/B5, 0)"
    ws_dash["H5"].font = FONT_KPI_PCT_VAL
    ws_dash["H5"].alignment = ALIGN_CENTER
    ws_dash["H5"].fill = FILL_KPI_PCT
    ws_dash["H5"].number_format = FORMAT_PERCENT

    # Bordes de las tarjetas KPI
    border_kpi_ingreso = Border(
        left=Side(style="thin", color="7DD3FC"),
        right=Side(style="thin", color="7DD3FC"),
        top=Side(style="thin", color="7DD3FC"),
        bottom=Side(style="thin", color="7DD3FC")
    )
    border_kpi_gasto = Border(
        left=Side(style="thin", color="FCA5A5"),
        right=Side(style="thin", color="FCA5A5"),
        top=Side(style="thin", color="FCA5A5"),
        bottom=Side(style="thin", color="FCA5A5")
    )
    border_kpi_saldo = Border(
        left=Side(style="thin", color="86EFAC"),
        right=Side(style="thin", color="86EFAC"),
        top=Side(style="thin", color="86EFAC"),
        bottom=Side(style="thin", color="86EFAC")
    )
    border_kpi_pct = Border(
        left=Side(style="thin", color="FCD34D"),
        right=Side(style="thin", color="FCD34D"),
        top=Side(style="thin", color="FCD34D"),
        bottom=Side(style="thin", color="FCD34D")
    )
    
    for r in (4, 5):
        for c in (2, 3): ws_dash.cell(row=r, column=c).border = border_kpi_ingreso
        for c in (4, 5): ws_dash.cell(row=r, column=c).border = border_kpi_gasto
        for c in (6, 7): ws_dash.cell(row=r, column=c).border = border_kpi_saldo
        for c in (8, 9): ws_dash.cell(row=r, column=c).border = border_kpi_pct

    ws_dash.row_dimensions[6].height = 14

    # ==========================================
    # SECCIONES INFERIORES: TABLA DE CATEGORÍAS Y RESUMEN PAGOS FIJOS
    # ==========================================
    ws_dash.merge_cells("A7:D7")
    ws_dash["A7"] = "📊 GASTOS POR CATEGORÍA (FÓRMULAS DINÁMICAS)"
    ws_dash["A7"].font = FONT_SECTION
    ws_dash["A7"].alignment = ALIGN_LEFT

    ws_dash.merge_cells("F7:I7")
    ws_dash["F7"] = "🔔 RESUMEN DE PAGOS FIJOS DEL MES"
    ws_dash["F7"].font = FONT_SECTION
    ws_dash["F7"].alignment = ALIGN_LEFT
    
    # Encabezados de Categorías
    ws_dash.row_dimensions[8].height = 24
    cat_headers = ["Categoría", "Total Gastado ($)", "% s/ Gastos", "% s/ Sueldo"]
    for idx, h in enumerate(cat_headers, start=1):
        c = ws_dash.cell(row=8, column=idx, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER_PRIMARY
        c.alignment = ALIGN_CENTER
        c.border = BORDER_HEADER

    # Encabezados de Resumen Pagos Fijos
    pf_headers = ["Concepto de Control", "Monto ($)", "Proporción", "Estado"]
    for idx, h in enumerate(pf_headers, start=6):
        c = ws_dash.cell(row=8, column=idx, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER_DARK
        c.alignment = ALIGN_CENTER
        c.border = BORDER_HEADER

    # Filas de Categorías
    lista_categorias = cat.get_categories_list()
    cat_start_row = 9
    for i, categoria in enumerate(lista_categorias):
        r = cat_start_row + i
        ws_dash.row_dimensions[r].height = 19
        fill_row = FILL_ZEBRA if i % 2 == 0 else FILL_WHITE
        
        # Categoría
        c_cat = ws_dash.cell(row=r, column=1, value=categoria)
        c_cat.alignment = ALIGN_LEFT
        c_cat.font = FONT_REGULAR
        c_cat.fill = fill_row
        c_cat.border = BORDER_CELL
        
        # Fórmula SUMIF para calcular el monto exacto desde la hoja de gastos
        c_monto = ws_dash.cell(row=r, column=2, value=f"=SUMIF('📝 Gastos del Mes'!$D$4:$D${last_gasto_row}, A{r}, '📝 Gastos del Mes'!$E$4:$E${last_gasto_row})")
        c_monto.alignment = ALIGN_RIGHT
        c_monto.font = FONT_REGULAR
        c_monto.number_format = FORMAT_CURRENCY
        c_monto.fill = fill_row
        c_monto.border = BORDER_CELL
        
        # % sobre gastos
        c_pct_g = ws_dash.cell(row=r, column=3, value=f"=IF($D$5>0, B{r}/$D$5, 0)")
        c_pct_g.alignment = ALIGN_RIGHT
        c_pct_g.font = FONT_REGULAR
        c_pct_g.number_format = FORMAT_PERCENT
        c_pct_g.fill = fill_row
        c_pct_g.border = BORDER_CELL

        # % sobre sueldo
        c_pct_s = ws_dash.cell(row=r, column=4, value=f"=IF($B$5>0, B{r}/$B$5, 0)")
        c_pct_s.alignment = ALIGN_RIGHT
        c_pct_s.font = FONT_REGULAR
        c_pct_s.number_format = FORMAT_PERCENT
        c_pct_s.fill = fill_row
        c_pct_s.border = BORDER_CELL

    cat_last_row = cat_start_row + len(lista_categorias) - 1
    cat_total_row = cat_last_row + 1
    
    # Total Categorías
    ws_dash.row_dimensions[cat_total_row].height = 22
    ws_dash.cell(row=cat_total_row, column=1, value="TOTAL EN CATEGORÍAS").font = FONT_TOTAL
    ws_dash.cell(row=cat_total_row, column=1).alignment = Alignment(horizontal="right", vertical="center")
    ws_dash.cell(row=cat_total_row, column=1).fill = FILL_TOTAL
    
    c_tot_cat_val = ws_dash.cell(row=cat_total_row, column=2, value=f"=SUM(B{cat_start_row}:B{cat_last_row})")
    c_tot_cat_val.font = FONT_TOTAL
    c_tot_cat_val.alignment = ALIGN_RIGHT
    c_tot_cat_val.number_format = FORMAT_CURRENCY
    c_tot_cat_val.fill = FILL_TOTAL
    
    c_tot_cat_pctg = ws_dash.cell(row=cat_total_row, column=3, value=f"=SUM(C{cat_start_row}:C{cat_last_row})")
    c_tot_cat_pctg.font = FONT_TOTAL
    c_tot_cat_pctg.alignment = ALIGN_RIGHT
    c_tot_cat_pctg.number_format = FORMAT_PERCENT
    c_tot_cat_pctg.fill = FILL_TOTAL

    c_tot_cat_pcts = ws_dash.cell(row=cat_total_row, column=4, value=f"=SUM(D{cat_start_row}:D{cat_last_row})")
    c_tot_cat_pcts.font = FONT_TOTAL
    c_tot_cat_pcts.alignment = ALIGN_RIGHT
    c_tot_cat_pcts.number_format = FORMAT_PERCENT
    c_tot_cat_pcts.fill = FILL_TOTAL
    
    for c_idx in range(1, 5):
        ws_dash.cell(row=cat_total_row, column=c_idx).border = BORDER_TOTAL

    # Filas de Resumen de Pagos Fijos (Fórmulas vinculadas a Hoja 3)
    pf_resumen = [
        ("1. Total Compromiso en Pagos Fijos", f"='🔔 Pagos Fijos'!C{total_pf_row}", "Presupuesto fijo del mes", "PREVISTO"),
        ("2. Pagos Fijos ya Abonados", f"=SUMIF('🔔 Pagos Fijos'!$F$4:$F${last_pf_row}, \"PAGADO\", '🔔 Pagos Fijos'!$C$4:$C${last_pf_row})", "Ya restados de la cuenta", "AL DÍA"),
        ("3. Pagos Fijos aún Pendientes", f"=G9-G10", "Restan abonar en el mes", "PENDIENTE"),
        ("4. % de Pagos Fijos Cumplidos", f"=IF(G9>0, G10/G9, 0)", "Avance de pagos fijos", "RATIO")
    ]
    
    for idx, (concepto, formula, detalle, estado_tag) in enumerate(pf_resumen, start=9):
        ws_dash.row_dimensions[idx].height = 20
        fill_pf_row = FILL_ZEBRA if idx % 2 == 0 else FILL_WHITE
        
        c_con = ws_dash.cell(row=idx, column=6, value=concepto)
        c_con.font = FONT_BOLD if "Total" in concepto else FONT_REGULAR
        c_con.alignment = ALIGN_LEFT
        c_con.fill = fill_pf_row
        c_con.border = BORDER_CELL
        
        c_val = ws_dash.cell(row=idx, column=7, value=formula)
        c_val.font = FONT_BOLD
        c_val.alignment = ALIGN_RIGHT
        c_val.number_format = FORMAT_PERCENT if "%" in concepto else FORMAT_CURRENCY
        c_val.fill = fill_pf_row
        c_val.border = BORDER_CELL
        
        c_det = ws_dash.cell(row=idx, column=8, value=detalle)
        c_det.font = FONT_SUBTITLE
        c_det.alignment = ALIGN_LEFT
        c_det.fill = fill_pf_row
        c_det.border = BORDER_CELL
        
        c_tag = ws_dash.cell(row=idx, column=9, value=estado_tag)
        c_tag.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color="BE123C")
        c_tag.alignment = ALIGN_CENTER
        c_tag.fill = fill_pf_row
        c_tag.border = BORDER_CELL

    # ==========================================
    # GRÁFICO CIRCULAR NATIVO (PIECHART)
    # ==========================================
    try:
        pie = PieChart()
        pie.title = f"Distribución de Gastos ({mes_nombre} {anio})"
        data = Reference(ws_dash, min_col=2, min_row=8, max_row=cat_last_row)
        labels = Reference(ws_dash, min_col=1, min_row=9, max_row=cat_last_row)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.width = 16
        pie.height = 10.5
        ws_dash.add_chart(pie, "F14")
    except Exception as e:
        print(f"Error creando grafico en Excel: {e}")


def _autoajustar_columnas(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignorar celdas combinadas de los banners para no ensanchar indebidamente
            if cell.row in (1, 2) and col_letter in ("A", "B", "C", "D", "E", "F", "G", "H", "I"):
                continue
            val_str = str(cell.value or "")
            if cell.value is not None:
                if str(cell.value).startswith("="):
                    val_str = "123,456.78"
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)
